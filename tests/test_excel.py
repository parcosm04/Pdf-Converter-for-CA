import os
import pytest
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook

from models import Transaction, ExceptionRecord, ValidationSummary
from excel_writer import write_excel_workbook

def test_write_excel_workbook(tmp_path):
    output_file = str(tmp_path / "test_statement.xlsx")
    
    transactions = [
        Transaction(
            date=date(2026, 1, 1),
            narration="SALARY PAYMENT",
            reference_number="REF12345",
            value_date=date(2026, 1, 1),
            debit=None,
            credit=Decimal("5000.00"),
            closing_balance=Decimal("5000.00"),
            page=1
        ),
        Transaction(
            date=date(2026, 1, 2),
            narration="RENT WITHDRAWAL",
            reference_number="REF9999",
            value_date=date(2026, 1, 2),
            debit=Decimal("1500.00"),
            credit=None,
            closing_balance=Decimal("3500.00"),
            page=1
        )
    ]
    
    summary = ValidationSummary(
        pages_processed=1,
        transactions_extracted=2,
        opening_balance=Decimal("0.00"),
        calculated_closing_balance=Decimal("3500.00"),
        pdf_closing_balance=Decimal("3500.00"),
        balance_match=True,
        total_debit=Decimal("1500.00"),
        total_credit=Decimal("5000.00"),
        duplicate_rows=0,
        missing_rows=0,
        broken_narrations=0,
        validation_status="PASS"
    )
    
    exceptions = [
        ExceptionRecord(
            page=1,
            transaction="Context details",
            issue="Some minor non-blocking issue",
            expected="N/A",
            found="N/A"
        )
    ]
    
    write_excel_workbook(output_file, transactions, summary, exceptions)
    
    # Assert file exists
    assert os.path.exists(output_file)
    
    # Load and check sheets
    wb = load_workbook(output_file)
    assert "Transactions" in wb.sheetnames
    assert "Validation Report" in wb.sheetnames
    assert "Exceptions" in wb.sheetnames
    
    # Check transactions sheet values
    ws_tx = wb["Transactions"]
    assert ws_tx["A2"].value.date() == date(2026, 1, 1)
    assert ws_tx["B2"].value == "SALARY PAYMENT"
    assert ws_tx["E2"].value is None
    assert float(ws_tx["F2"].value) == 5000.00
    
    # Check validation report values
    ws_val = wb["Validation Report"]
    assert ws_val["B3"].value == "PASS"
    assert int(ws_val["B6"].value) == 2  # Transactions Extracted
    assert float(ws_val["B8"].value) == 1500.00  # Total Debit
    assert float(ws_val["B9"].value) == 5000.00  # Total Credit
    assert float(ws_val["B7"].value) == 0.00     # Opening Balance
    
    # Check exceptions sheet values
    ws_ex = wb["Exceptions"]
    assert ws_ex["C2"].value == "Some minor non-blocking issue"
