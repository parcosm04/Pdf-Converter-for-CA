import os
from typing import List
from datetime import date
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import config
from models import Transaction, ExceptionRecord, ValidationSummary
from loguru import logger

def write_excel_workbook(
    file_path: str,
    transactions: List[Transaction],
    validation_summary: ValidationSummary,
    exceptions: List[ExceptionRecord]
) -> None:
    """
    Writes transactions, validation report, and exceptions into a professional Excel file.
    """
    logger.info(f"Writing parsed results to Excel: {file_path}")
    wb = Workbook()
    
    # Remove default sheet
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)
        
    # Styles Definition
    font_family = config.EXCEL_FONT_NAME
    
    font_header = Font(name=font_family, size=11, bold=True, color=config.EXCEL_HEADER_FG)
    fill_header = PatternFill(start_color=config.EXCEL_HEADER_BG, end_color=config.EXCEL_HEADER_BG, fill_type="solid")
    
    font_data = Font(name=font_family, size=10, color="000000")
    font_bold = Font(name=font_family, size=10, bold=True, color="000000")
    
    fill_alt = PatternFill(start_color=config.EXCEL_ROW_ALT_BG, end_color=config.EXCEL_ROW_ALT_BG, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Border styles
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    border_double_bottom = Border(bottom=Side(border_style="double", color="1B365D"), top=Side(border_style="thin", color="D3D3D3"))
    
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    
    # Number formats
    format_currency = '"$"#,##0.00;("$"#,##0.00);"-"'
    format_date = "YYYY-MM-DD"
    
    # -------------------------------------------------------------
    # SHEET 1: Transactions
    # -------------------------------------------------------------
    ws_tx = wb.create_sheet(title="Transactions")
    ws_tx.views.sheetView[0].showGridLines = True
    
    headers_tx = ["Date", "Narration", "Reference Number", "Value Date", "Debit", "Credit", "Closing Balance", "Page"]
    ws_tx.append(headers_tx)
    
    # Format header row
    for col_num in range(1, len(headers_tx) + 1):
        cell = ws_tx.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        
    # Write data rows
    for r_idx, tx in enumerate(transactions, start=2):
        row_data = [
            tx.date,
            tx.narration,
            tx.reference_number or "-",
            tx.value_date or tx.date,
            tx.debit,
            tx.credit,
            tx.closing_balance,
            tx.page
        ]
        ws_tx.append(row_data)
        
        # Decide row fill color
        row_fill = fill_alt if r_idx % 2 == 0 else fill_white
        
        # Formatting individual cells
        for c_idx in range(1, len(row_data) + 1):
            cell = ws_tx.cell(row=r_idx, column=c_idx)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = border_data
            
            # Alignments & Number formats
            if c_idx in [1, 4]:  # Dates
                cell.number_format = format_date
                cell.alignment = align_center
            elif c_idx in [2, 3]:  # Narration, Ref
                cell.alignment = align_left
            elif c_idx in [5, 6, 7]:  # Debit, Credit, Balance
                cell.number_format = format_currency
                cell.alignment = align_right
            elif c_idx == 8:  # Page
                cell.alignment = align_center
                
    # Freeze Header row
    ws_tx.freeze_panes = "A2"
    
    # Enable Autofilter
    if len(transactions) > 0:
        ws_tx.auto_filter.ref = f"A1:H{len(transactions) + 1}"
        
    # Auto-fit columns
    for col in ws_tx.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Check length of values to adjust width
            val = cell.value
            if isinstance(val, date):
                val_str = val.strftime("%Y-%m-%d")
            elif isinstance(val, (int, float, Decimal)):
                val_str = f"{val:,.2f}"
            else:
                val_str = str(val or "")
            max_len = max(max_len, len(val_str))
        ws_tx.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # -------------------------------------------------------------
    # SHEET 2: Validation Report
    # -------------------------------------------------------------
    ws_val = wb.create_sheet(title="Validation Report")
    ws_val.views.sheetView[0].showGridLines = True
    
    # Styling values specifically for reports
    fill_pass = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")  # Soft Green
    font_pass = Font(name=font_family, size=11, bold=True, color="155724")
    
    fill_fail = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")  # Soft Red
    font_fail = Font(name=font_family, size=11, bold=True, color="721C24")
    
    # Title
    ws_val.merge_cells("A1:C1")
    title_cell = ws_val["A1"]
    title_cell.value = "BANK STATEMENT VALIDATION REPORT"
    title_cell.font = Font(name=font_family, size=14, bold=True, color="1B365D")
    title_cell.alignment = align_left
    
    # Validation status banner
    status_row = 3
    ws_val.cell(row=status_row, column=1, value="Overall Validation Status").font = font_bold
    status_cell = ws_val.cell(row=status_row, column=2, value=validation_summary.validation_status)
    status_cell.alignment = align_center
    if validation_summary.validation_status == "PASS":
        status_cell.fill = fill_pass
        status_cell.font = font_pass
    else:
        status_cell.fill = fill_fail
        status_cell.font = font_fail
        
    # Metrics
    metrics = [
        ("Pages Processed", validation_summary.pages_processed, None),
        ("Transactions Extracted", validation_summary.transactions_extracted, None),
        ("Opening Balance", validation_summary.opening_balance, format_currency),
        ("Total Debit (-)", validation_summary.total_debit, format_currency),
        ("Total Credit (+)", validation_summary.total_credit, format_currency),
        ("Calculated Closing Balance", validation_summary.calculated_closing_balance, format_currency),
        ("PDF Statement Closing Balance", validation_summary.pdf_closing_balance, format_currency),
        ("Balances Match (Calculated vs PDF)", "YES" if validation_summary.balance_match else "NO", None),
        ("Duplicate Rows Detected", validation_summary.duplicate_rows, None),
        ("Balance Discrepancies", validation_summary.missing_rows, None),
        ("Empty/Broken Narrations", validation_summary.broken_narrations, None)
    ]
    
    curr_row = 5
    for label, val, num_fmt in metrics:
        cell_lbl = ws_val.cell(row=curr_row, column=1, value=label)
        cell_lbl.font = font_data
        cell_lbl.border = border_data
        
        cell_val = ws_val.cell(row=curr_row, column=2, value=val)
        cell_val.font = font_bold
        cell_val.border = border_data
        
        # Color yes/no
        if label.startswith("Balances Match"):
            cell_val.alignment = align_center
            if val == "YES":
                cell_val.fill = fill_pass
                cell_val.font = font_pass
            else:
                cell_val.fill = fill_fail
                cell_val.font = font_fail
        # Highlight errors
        elif label in ["Duplicate Rows Detected", "Balance Discrepancies", "Empty/Broken Narrations"]:
            cell_val.alignment = align_center
            if val > 0:
                cell_val.fill = fill_fail
                cell_val.font = font_fail
            else:
                cell_val.fill = fill_pass
                cell_val.font = font_pass
        else:
            if isinstance(val, (int, float, Decimal)):
                cell_val.alignment = align_right
                if num_fmt:
                    cell_val.number_format = num_fmt
            else:
                cell_val.alignment = align_left
                
        curr_row += 1
        
    ws_val.column_dimensions["A"].width = 38
    ws_val.column_dimensions["B"].width = 20
    
    # -------------------------------------------------------------
    # SHEET 3: Exceptions
    # -------------------------------------------------------------
    ws_ex = wb.create_sheet(title="Exceptions")
    ws_ex.views.sheetView[0].showGridLines = True
    
    headers_ex = ["Page", "Transaction/Context", "Issue Description", "Expected Value", "Found Value"]
    ws_ex.append(headers_ex)
    
    for col_num in range(1, len(headers_ex) + 1):
        cell = ws_ex.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        
    if not exceptions:
        ws_ex.append(["-", "No exceptions or errors were detected during parsing and verification.", "-", "-", "-"])
        # Format the empty status line
        for c_idx in range(1, 6):
            cell = ws_ex.cell(row=2, column=c_idx)
            cell.font = Font(name=font_family, size=10, italic=True)
            cell.alignment = align_left if c_idx == 2 else align_center
            cell.border = border_data
    else:
        for r_idx, ex in enumerate(exceptions, start=2):
            row_data = [
                ex.page,
                ex.transaction,
                ex.issue,
                ex.expected or "-",
                ex.found or "-"
            ]
            ws_ex.append(row_data)
            
            # Format row
            row_fill = fill_alt if r_idx % 2 == 0 else fill_white
            for c_idx in range(1, len(row_data) + 1):
                cell = ws_ex.cell(row=r_idx, column=c_idx)
                cell.font = font_data
                cell.fill = row_fill
                cell.border = border_data
                
                if c_idx in [1]:
                    cell.alignment = align_center
                elif c_idx in [2, 3]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center
                    
    # Auto-fit columns for exceptions sheet
    for col in ws_ex.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            max_len = max(max_len, len(str(cell.value or "")))
        ws_ex.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)  # Cap width at 50 to avoid overly wide columns
        
    # Ensure folder path exists
    os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
    
    # Save Workbook
    wb.save(file_path)
    logger.info(f"Excel workbook saved successfully with {len(transactions)} transactions.")
